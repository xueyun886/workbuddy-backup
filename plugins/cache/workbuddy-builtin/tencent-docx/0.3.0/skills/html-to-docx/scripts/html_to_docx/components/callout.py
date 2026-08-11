"""Callout component renderer.

Renders ``data-component="callout"`` elements as a 1×1 table with a
colored left border and light background, keyed by ``data-variant``.
"""
from __future__ import annotations

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx.document import Document as DocxDocument
from docx.text.paragraph import Paragraph
from bs4 import Tag

from . import register
from ..style_mapper import apply_run_styles, apply_paragraph_styles, _parse_style, _parse_color
from ._cell_richtext import render_children_into_cell

_VARIANT_COLORS: dict[str, tuple[str, str]] = {
    # (border-left-color, background-color)
    # background: rgba(r,g,b, 0.06) on white ≈ blended hex
    "info":    ("01579B", "F0F5FA"),  # rgba(1,87,155, 0.06) on #FFF
    "warning": ("E65100", "FDF3EF"),  # rgba(230,81,0, 0.06) on #FFF
    "danger":  ("C62828", "FBF0F0"),  # rgba(198,40,40, 0.06) on #FFF
    "success": ("2E7D32", "F0F7F0"),  # rgba(46,125,50, 0.06) on #FFF
}


@register("callout")
def render_callout(element: Tag, document: DocxDocument, anchor: Paragraph | None ) -> None:
    """Render a callout component into *document*."""
    variant = element.get("data-variant", "info").lower()
    border_hex, bg_hex = _VARIANT_COLORS.get(variant, _VARIANT_COLORS["info"])

    # Locate the rich-text content element (fall back to the callout itself).
    content_el = element.find(class_="callout-content") or element

    # 1-row 1-col table
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)

    # Background shading + border color — inline style overrides variant default
    style_attr = element.get("style", "")
    if style_attr:
        props = _parse_style(style_attr)
        if "background-color" in props or "background" in props:
            bg_val = props.get("background-color") or props.get("background", "")
            color = _parse_color(bg_val)
            if color is not None:
                bg_hex = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"
        if "border-left-color" in props:
            color = _parse_color(props["border-left-color"])
            if color is not None:
                border_hex = f"{color[0]:02X}{color[1]:02X}{color[2]:02X}"

    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), bg_hex)
    tcPr.append(shd)

    # Left border with variant color
    tcBorders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")   # 3pt thick
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), border_hex)
    tcBorders.append(left)
    tcPr.append(tcBorders)

    # Cell content — render rich text (preserve <strong>/multi <p>/<li>),
    # replacing the legacy get_text(separator="\n") + para.text flattening.
    render_children_into_cell(content_el, cell)

    # Apply element-level inline style to all rendered paragraphs/runs.
    style_attr = element.get("style", "")
    if style_attr:
        for para in cell.paragraphs:
            apply_paragraph_styles(para, style_attr)
            for run in para.runs:
                apply_run_styles(run, style_attr)
    if not anchor is None:
        anchor._element.addprevious(table._element)
        anchor._element.getparent().remove(anchor._element)
